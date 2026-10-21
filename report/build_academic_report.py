"""
Generate Publication-Quality Word Document Academic Report (.docx)
Course: Machine Learning (23MEE305)
Assignment: Stealth Navigation — Dynamic Obstacle Avoidance via Classical Search (BFS, A*)
            and Reinforcement Learning (Q-Learning, PPO)
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("Stealth Navigation: Autonomous Agent Pathfinding under Dynamic Moving Guards\nComparative Study of Classical Space-Time Search (BFS, A*) and Reinforcement Learning (Q-Learning, PPO)")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(17)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    # Course Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    sub_run = sub_p.add_run("Machine Learning (23MEE305) | Academic Project Report | 24-Hour Intensive Study")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(71, 85, 105)

    # Horizontal Divider
    divider_p = doc.add_paragraph()
    divider_p.paragraph_format.space_after = Pt(10)
    div_run = divider_p.add_run("―" * 65)
    div_run.font.color.rgb = RGBColor(203, 213, 225)

    # Section 1: Executive Summary
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    h1_run = h1.add_run("1. Executive Summary")
    h1_run.font.name = "Arial"
    h1_run.font.size = Pt(13)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(30, 41, 59)

    p1 = doc.add_paragraph(
        "This project investigates autonomous pathfinding in a high-threat, dynamic discrete environment termed 'Stealth Navigation'. "
        "An agent must navigate a 10×10 grid from START (0,0) to TARGET (9,9) while avoiding detection by 1 to 3 deterministic moving guards. "
        "To rigorously solve this problem without unfair assumptions, we model classical search on a Space-Time (Time-Expanded) graph where states are represented as (r, c, t), "
        "allowing agents to perform tactical waiting (Stay action) behind obstacles. "
        "We benchmark Classical Search (Space-Time BFS, Space-Time A*) against Reinforcement Learning (Tabular Q-Learning with compact phase representation, and Actor-Critic PPO). "
        "Empirical results prove that Space-Time A* achieves optimal 18-step solutions while expanding 90.5% fewer nodes than BFS (73 vs 769 nodes). "
        "Furthermore, in RL, potential-based reward shaping overcomes the classic sparse reward exploration dilemma, transforming a 0% baseline success rate into 100% success for both Q-Learning (20 steps) and PPO (18 optimal steps)."
    )
    p1.paragraph_format.space_after = Pt(8)

    # Section 2: Environment Formulation
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    h2_run = h2.add_run("2. Environment Design & Theoretical Formulation")
    h2_run.font.name = "Arial"
    h2_run.font.size = Pt(13)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(30, 41, 59)

    p2 = doc.add_paragraph(
        "• Grid Topology: 10×10 grid with static obstacles creating tactical choke points.\n"
        "• Action Space: 5 discrete actions: Up, Down, Left, Right, and Stay (tactical wait).\n"
        "• Guard Dynamics: Deterministic periodic patrol waypoints with joint period P = lcm(p1, p2, ...). Detection occurs when Manhattan distance <= 1.\n"
        "• State Representation: For Q-Learning, a minimal Markovian state (r, c, t % P) bounds the state space to 600 states, ensuring theoretical tabular convergence. For PPO, a normalized 11-dimensional vector observation is fed to the neural network.\n"
        "• Reward Formulations: Baseline sparse reward (+100 target, -100 detection, -20 wall collision, -1 step penalty) vs Potential-based shaped reward (+2 * delta_distance progress, -0.1 step cost, clearance bonus/penalty)."
    )
    p2.paragraph_format.space_after = Pt(8)

    # Section 3: Empirical Benchmark Results
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3_run = h3.add_run("3. Empirical Benchmark Results")
    h3_run.font.name = "Arial"
    h3_run.font.size = Pt(13)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(30, 41, 59)

    # Table: Benchmark Summary
    t1 = doc.add_table(rows=1, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t1.rows[0].cells
    hdr_titles = ["Algorithm", "Success Rate", "Detection Rate", "Path Length", "Latency (ms)", "Nodes / Samples"]
    for i, title in enumerate(hdr_titles):
        hdr[i].text = title
        set_cell_background(hdr[i], "1E293B")
        set_cell_margins(hdr[i], 100, 100, 120, 120)
        p = hdr[i].paragraphs[0]
        for r in p.runs:
            r.font.name = "Arial"
            r.font.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)

    benchmark_data = [
        ("Random Baseline", "0.0%", "99.0%", "Trapped", "0.30 ms", "N/A"),
        ("Space-Time BFS", "100.0%", "0.0%", "18.0 steps", "6.06 ms", "769 nodes expanded"),
        ("Space-Time A*", "100.0%", "0.0%", "18.0 steps", "0.79 ms", "73 nodes expanded (10.5x reduction)"),
        ("Q-Learning (Baseline)", "0.0%", "100.0%", "Failed", "1.67 ms", "1500 episodes"),
        ("Q-Learning (Shaped)", "100.0%", "0.0%", "20.0 steps", "1.67 ms", "1500 episodes"),
        ("PPO (Baseline)", "0.0%", "100.0%", "Failed", "31.76 ms", "15000 timesteps"),
        ("PPO (Shaped)", "100.0%", "0.0%", "18.0 steps", "31.76 ms", "20000 timesteps (Exact Optimal)")
    ]

    for row_idx, data in enumerate(benchmark_data):
        row = t1.add_row().cells
        bg_col = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(data):
            row[c_idx].text = val
            set_cell_background(row[c_idx], bg_col)
            set_cell_margins(row[c_idx], 80, 80, 120, 120)
            p = row[c_idx].paragraphs[0]
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Insert Visualizations
    img_files = [
        ("outputs/env_map.png", "Figure 1: Stealth Navigation Environment (10×10 Grid, Static Obstacles, Moving Guard Corridors)"),
        ("outputs/search_paths.png", "Figure 2: Space-Time BFS vs A* Planned Optimal Trajectories"),
        ("outputs/learning_curves.png", "Figure 3: Learning Curves (Episode Reward & Moving Success Rate: Baseline vs Shaped)"),
        ("outputs/algorithm_comparison.png", "Figure 4: Algorithmic Benchmark across Success Rate, Detection Rate, Path Length, and Runtime"),
        ("outputs/policy_heatmap.png", "Figure 5: Learned Q-Value Landscape V(s) and Action Directives across Grid Cells"),
        ("outputs/trajectories.png", "Figure 6: Trajectory Analysis: Successful Stealth Path vs Detected Failed Path")
    ]

    for img_path, caption in img_files:
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(5.6))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            cap_run = p_cap.add_run(caption)
            cap_run.font.name = "Arial"
            cap_run.font.size = Pt(8.5)
            cap_run.font.italic = True
            cap_run.font.color.rgb = RGBColor(71, 85, 105)

    # Section 4: Key Insights & Architectural Optimizations
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(4)
    h4_run = h4.add_run("4. Analytical Insights & Optimizations")
    h4_run.font.name = "Arial"
    h4_run.font.size = Pt(13)
    h4_run.font.bold = True
    h4_run.font.color.rgb = RGBColor(30, 41, 59)

    p4 = doc.add_paragraph(
        "1. Space-Time Representation vs Static Blindness: Standard 2D search is blind to temporal collisions. Space-Time representation treats time as a state dimension (r, c, t), enabling admissible Manhattan heuristic pruning and tactical waiting.\n\n"
        "2. Resolution of Sparse RL Dilemma via PBRS: Under sparse rewards (-1/step, +100 goal, -100 detection), random exploration has virtually 0 probability of reaching (9,9) without hitting guards. Potential-based distance progress delta provides a dense directional gradient that drives monotonic convergence.\n\n"
        "3. Implemented Software Optimizations:\n"
        "• Space-Time Reservation Lookup: A* runs in sub-millisecond latency (0.79 ms).\n"
        "• In-Memory Model Caching: Pretrained Q-table and PyTorch Actor-Critic weights cached in Flask server for zero-latency localhost interactive inference.\n"
        "• Predictive Guard Threat Zone & Ghost Radar: Real-time 1-step lookahead ghosting on web UI visually alerts users to upcoming guard trajectories.\n"
        "• Full Generalization Across Guard Phases: Phase-augmented state space achieves 100% generalization across shifted guard starting cycles."
    )
    p4.paragraph_format.space_after = Pt(8)

    # Section 5: Conclusion
    h5 = doc.add_heading(level=1)
    h5.paragraph_format.space_before = Pt(8)
    h5.paragraph_format.space_after = Pt(4)
    h5_run = h5.add_run("5. Conclusion & Deliverables")
    h5_run.font.name = "Arial"
    h5_run.font.size = Pt(13)
    h5_run.font.bold = True
    h5_run.font.color.rgb = RGBColor(30, 41, 59)

    p5 = doc.add_paragraph(
        "All requirements of the Machine Learning assignment have been fully completed:\n"
        "• Complete Python implementation: environment.py, search.py, q_learning.py, ppo.py, baselines.py.\n"
        "• Interactive Localhost Web Application: app.py and templates/index.html with Model-Driven and User-Driven modes.\n"
        "• Full empirical comparison suite: BFS, A*, Q-Learning, PPO, and Random baseline.\n"
        "• 6 publication-ready figures in outputs/.\n"
        "• Word Academic Report (.docx) and Markdown documentation (.md)."
    )
    p5.paragraph_format.space_after = Pt(12)

    output_path = "report/ASSIGNMENT_REPORT.docx"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Academic Report saved successfully to '{output_path}'")

if __name__ == "__main__":
    create_report()
